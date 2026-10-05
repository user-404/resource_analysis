from types import SimpleNamespace
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.main
from app.db import Base
from app.models import CollectionRun


def test_collection_can_be_queued_and_polled(monkeypatch):
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    monkeypatch.setattr(app.main, "SessionLocal", sessionmaker(bind=engine, expire_on_commit=False))
    monkeypatch.setattr(app.main, "initialize_database", lambda: Base.metadata.create_all(engine))
    monkeypatch.setattr(app.main, "authorize", lambda authorization: None)
    monkeypatch.setattr(app.main, "get_settings", lambda: SimpleNamespace(environment="development"))

    payload = {
        "environment": "staging",
        "namespace": "payments",
        "daysToCollect": 3,
        "kubeconfigSecretRef": "env://KUBECONFIG_TEST",
        "prometheusUrl": "https://prometheus.example.test",
        "prometheusSecretRef": "env://PROMETHEUS_TEST",
    }
    with TestClient(app.main.app) as client:
        dashboard = client.get("/")
        assert dashboard.status_code == 200
        assert "New collection" in dashboard.text
        assert "position: fixed; top: 0; left: 0; right: 0" in dashboard.text
        assert "padding-top: 94px" in dashboard.text
        assert "Filter workloads by one or more names" in dashboard.text
        assert "CronJob results include usage from their associated Job runs." in dashboard.text
        assert dashboard.headers["cache-control"] == "no-store"
        invalid_payload = {**payload, "prometheusUrl": "https://user:password@prometheus.example.test"}
        assert client.post("/v1/collections", json=invalid_payload).status_code == 422
        legacy_identifier_payload = {
            **payload,
            "userId": "user-1",
            "organizationId": "org-1",
            "configId": "config-1",
        }
        assert client.post("/v1/collections", json=legacy_identifier_payload).status_code == 422
        response = client.post("/v1/collections", json=payload)
        assert response.status_code == 202
        run_id = response.json()["runId"]
        no_prometheus_auth_payload = {
            key: value for key, value in payload.items() if key != "prometheusSecretRef"
        }
        no_auth_response = client.post("/v1/collections", json=no_prometheus_auth_payload)
        assert no_auth_response.status_code == 202, no_auth_response.text

        status = client.get(f"/v1/collections/{run_id}")
        assert status.status_code == 200
        assert status.json()["status"] == "queued"
        assert not {"userId", "organizationId", "configId"} & status.json().keys()

        history = client.get("/v1/collections")
        assert history.status_code == 200
        assert {entry["runId"] for entry in history.json()} >= {
            run_id,
            no_auth_response.json()["runId"],
        }

        demo_response = client.post(
            "/v1/demo/collections",
            json={
                "namespace": "payments",
                "daysToCollect": 1,
                "currency": "EUR",
                "cpuRatePerCoreHour": 0.04,
                "memoryRatePerGiBHour": 0.005,
            },
        )
        assert demo_response.status_code == 201, demo_response.text
        demo_run_id = demo_response.json()["runId"]
        demo_status = client.get(f"/v1/collections/{demo_run_id}")
        assert demo_status.status_code == 200
        assert demo_status.json()["status"] == "completed"
        assert demo_status.json()["environment"] == "DEMO (synthetic)"
        demo_aggregates = client.get(f"/v1/collections/{demo_run_id}/aggregates")
        assert demo_aggregates.status_code == 200
        assert len(demo_aggregates.json()) == 5
        assert any(row["workloadType"] == "CronJob" for row in demo_aggregates.json())
        raw_page = client.get(
            f"/v1/collections/{demo_run_id}/raw-samples",
            params={"workloadName": "checkout", "limit": 2},
        )
        assert raw_page.status_code == 200, raw_page.text
        assert raw_page.json()["workloadType"] == "Deployment"
        assert raw_page.json()["total"] == 24
        assert len(raw_page.json()["items"]) == 2
        assert {
            "podName", "sampledAt", "cpuCores", "memoryBytes",
            "allocatedCpuRequestCores", "allocatedMemoryLimitBytes",
        } <= raw_page.json()["items"][0].keys()
        second_page = client.get(
            f"/v1/collections/{demo_run_id}/raw-samples",
            params={"workloadName": "checkout", "offset": 2, "limit": 2},
        )
        assert second_page.status_code == 200
        assert second_page.json()["items"][0]["sampledAt"] != raw_page.json()["items"][0]["sampledAt"]
        assert client.get(
            f"/v1/collections/{demo_run_id}/raw-samples",
            params={"workloadName": "missing-workload"},
        ).status_code == 404
        export_response = client.get(f"/v1/collections/{demo_run_id}/export.xlsx")
        assert export_response.status_code == 200, export_response.text
        assert export_response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert export_response.headers["content-disposition"].endswith(f'"resource-analysis-{demo_run_id}.xlsx"')
        workbook = load_workbook(BytesIO(export_response.content), read_only=True, data_only=True)
        assert workbook.sheetnames == ["Workload Summary", "Raw Samples 1"]
        summary_sheet = workbook["Workload Summary"]
        summary_rows = list(summary_sheet.values)
        assert len(summary_rows) == 6
        assert any(row[0] == "checkout" for row in summary_rows[1:])
        raw_sheet = workbook["Raw Samples 1"]
        raw_rows = list(raw_sheet.values)
        assert len(raw_rows) == 121
        assert any(row[0] == "checkout" and row[3].startswith("checkout-") for row in raw_rows[1:])
        workbook.close()
        assert client.get(f"/v1/collections/{run_id}/export.xlsx").status_code == 409
        cost_estimate = demo_aggregates.json()[0]["costEstimate"]
        assert cost_estimate["currency"] == "EUR"
        assert cost_estimate["monthlyHours"] == 730
        assert cost_estimate["yearlyHours"] == 8760
        assert abs(cost_estimate["savingsYearly"] - cost_estimate["savingsMonthly"] * 12) < 0.06

        no_rates_response = client.post(
            "/v1/demo/collections",
            json={"namespace": "payments", "daysToCollect": 1},
        )
        assert no_rates_response.status_code == 201
        no_rates = client.get(f"/v1/collections/{no_rates_response.json()['runId']}/aggregates")
        assert no_rates.status_code == 200
        assert all(row["costEstimate"] is None for row in no_rates.json())

    with sessionmaker(bind=engine)() as db:
        run = db.get(CollectionRun, run_id)
        assert run is not None
        assert run.input_json["kubeconfig_secret_ref"] == "env://KUBECONFIG_TEST"
        assert run.input_json["prometheus_secret_ref"] == "env://PROMETHEUS_TEST"
        no_auth_run = db.get(CollectionRun, no_auth_response.json()["runId"])
        assert no_auth_run is not None
        assert no_auth_run.input_json["prometheus_secret_ref"] is None
        demo_run = db.get(CollectionRun, demo_run_id)
        assert demo_run is not None
        assert demo_run.input_json["synthetic"] is True
        assert demo_run.input_json["cpu_rate_per_core_hour"] == 0.04
        assert "prometheus_secret_ref" not in demo_run.input_json
    engine.dispose()


def test_rate_inputs_must_be_provided_as_a_pair():
    from app.schemas import DemoCollectionCreate
    from pydantic import ValidationError

    try:
        DemoCollectionCreate(cpuRatePerCoreHour=0.05)
    except ValidationError as exc:
        assert "Provide both CPU and memory rates" in str(exc)
    else:
        raise AssertionError("One-sided cost rates should be rejected")
