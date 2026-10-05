from contextlib import asynccontextmanager
import json
from pathlib import Path
from tempfile import SpooledTemporaryFile
from typing import Annotated

from fastapi import Depends, FastAPI, Header, HTTPException, Query
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session
from starlette.background import BackgroundTask

from app.auth import authorize
from app.config import get_settings
from app.costs import estimate_cost
from app.demo import create_demo_collection
from app.db import SessionLocal, initialize_database
from app.models import CollectionRun, RawUsageSample, WorkloadAggregate
from app.schemas import (
    AggregateResponse,
    CollectionAccepted,
    CollectionCreate,
    CollectionStatus,
    DemoCollectionCreate,
    RawUsageSamplesPage,
)
from app.service import enqueue_collection


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


@asynccontextmanager
async def lifespan(_: FastAPI):
    initialize_database()
    yield


app = FastAPI(title="Kubernetes Resource Analysis", version="0.1.0", lifespan=lifespan)
STATIC_DIR = Path(__file__).parent / "static"


def _excel_text(value: str) -> str:
    if value.lstrip(" \t\r\n")[:1] in {"=", "+", "-", "@"}:
        return f"'{value}"
    return value


@app.get("/healthz")
def healthz():
    return {"status": "ok"}


@app.get("/", include_in_schema=False)
def dashboard():
    return FileResponse(STATIC_DIR / "index.html", headers={"Cache-Control": "no-store"})


@app.post("/v1/collections", response_model=CollectionAccepted, response_model_by_alias=True, status_code=202)
def create_collection(
    payload: CollectionCreate,
    db: Annotated[Session, Depends(get_db)],
    authorization: Annotated[str | None, Header()] = None,
):
    authorize(authorization)
    run = enqueue_collection(db, payload.model_dump(by_alias=False, mode="json"))
    return CollectionAccepted(run_id=run.id, status=run.status)


@app.post(
    "/v1/demo/collections",
    response_model=CollectionAccepted,
    response_model_by_alias=True,
    status_code=201,
)
def create_demo(
    payload: DemoCollectionCreate,
    db: Annotated[Session, Depends(get_db)],
    authorization: Annotated[str | None, Header()] = None,
):
    if get_settings().environment == "production":
        raise HTTPException(status_code=404, detail="Not found")
    authorize(authorization)
    run = create_demo_collection(
        db,
        payload.namespace,
        payload.days_to_collect,
        payload.currency,
        payload.cpu_rate_per_core_hour,
        payload.memory_rate_per_gib_hour,
    )
    return CollectionAccepted(run_id=run.id, status=run.status)


def get_run(run_id: str, db: Session) -> CollectionRun:
    run = db.get(CollectionRun, run_id)
    if run is None:
        raise HTTPException(status_code=404, detail="Collection run not found")
    return run


@app.get("/v1/collections/{run_id}", response_model=CollectionStatus, response_model_by_alias=True)
def collection_status(
    run_id: str,
    db: Annotated[Session, Depends(get_db)],
    authorization: Annotated[str | None, Header()] = None,
):
    authorize(authorization)
    return get_run(run_id, db)


@app.get("/v1/collections", response_model=list[CollectionStatus], response_model_by_alias=True)
def collection_history(
    db: Annotated[Session, Depends(get_db)],
    authorization: Annotated[str | None, Header()] = None,
    limit: Annotated[int, Query(ge=1, le=50)] = 20,
):
    authorize(authorization)
    return db.scalars(
        select(CollectionRun)
        .order_by(CollectionRun.created_at.desc())
        .limit(limit)
    ).all()


@app.get(
    "/v1/collections/{run_id}/aggregates",
    response_model=list[AggregateResponse],
    response_model_by_alias=True,
)
def collection_aggregates(
    run_id: str,
    db: Annotated[Session, Depends(get_db)],
    authorization: Annotated[str | None, Header()] = None,
):
    authorize(authorization)
    run = get_run(run_id, db)
    if run.status != "completed":
        raise HTTPException(status_code=409, detail=f"Collection is {run.status}")
    rows = db.scalars(
        select(WorkloadAggregate)
        .where(WorkloadAggregate.run_id == run_id)
        .order_by(WorkloadAggregate.workload_type, WorkloadAggregate.workload_name)
    ).all()
    results = []
    for row in rows:
        result = AggregateResponse.model_validate(row)
        result.cost_estimate = estimate_cost(row, run.input_json)
        results.append(result)
    return results


@app.get("/v1/collections/{run_id}/export.xlsx", include_in_schema=True)
def export_collection_excel(
    run_id: str,
    db: Annotated[Session, Depends(get_db)],
    authorization: Annotated[str | None, Header()] = None,
):
    authorize(authorization)
    run = get_run(run_id, db)
    if run.status != "completed":
        raise HTTPException(status_code=409, detail=f"Collection is {run.status}")

    workbook = Workbook(write_only=True)
    summary = workbook.create_sheet("Workload Summary")
    summary.append([
        "Workload", "Type", "Pods observed", "Time points", "Average CPU (cores)", "Peak CPU (cores)",
        "Average memory (bytes)", "Peak memory (bytes)", "CPU request (cores)", "CPU limit (cores)",
        "Suggested CPU request (cores)", "Suggested CPU limit (cores)", "Memory request (bytes)",
        "Memory limit (bytes)", "Suggested memory request (bytes)", "Suggested memory limit (bytes)",
        "Monthly savings", "Yearly savings", "Currency", "Recommendation policy",
    ])
    aggregates = db.scalars(
        select(WorkloadAggregate)
        .where(WorkloadAggregate.run_id == run_id)
        .order_by(WorkloadAggregate.workload_type, WorkloadAggregate.workload_name)
    )
    for row in aggregates:
        cost = estimate_cost(row, run.input_json)
        summary.append([
            _excel_text(row.workload_name), _excel_text(row.workload_type), row.pod_count, row.sample_count,
            row.average_cpu_cores, row.peak_cpu_cores, row.average_memory_bytes, row.peak_memory_bytes,
            row.allocated_cpu_request_cores, row.allocated_cpu_limit_cores,
            row.suggested_cpu_request_cores, row.suggested_cpu_limit_cores,
            row.allocated_memory_request_bytes, row.allocated_memory_limit_bytes,
            row.suggested_memory_request_bytes, row.suggested_memory_limit_bytes,
            cost.savings_monthly if cost else None, cost.savings_yearly if cost else None,
            cost.currency if cost else None, json.dumps(row.recommendation_policy, sort_keys=True),
        ])

    raw_headers = [
        "Workload", "Type", "Namespace", "Pod", "Sampled at", "CPU usage (cores)", "Memory usage (bytes)",
        "CPU request (cores)", "CPU limit (cores)", "Memory request (bytes)", "Memory limit (bytes)",
    ]
    raw_sheet_number = 1
    raw_sheet = workbook.create_sheet(f"Raw Samples {raw_sheet_number}")
    raw_sheet.append(raw_headers)
    samples = db.scalars(
        select(RawUsageSample)
        .where(RawUsageSample.run_id == run_id)
        .order_by(RawUsageSample.sampled_at, RawUsageSample.workload_name, RawUsageSample.pod_name)
        .execution_options(yield_per=1000)
    )
    rows_in_sheet = 0
    max_data_rows = 1_048_575
    for sample in samples:
        if rows_in_sheet >= max_data_rows:
            raw_sheet_number += 1
            raw_sheet = workbook.create_sheet(f"Raw Samples {raw_sheet_number}")
            raw_sheet.append(raw_headers)
            rows_in_sheet = 0
        raw_sheet.append([
            _excel_text(sample.workload_name), _excel_text(sample.workload_type),
            _excel_text(sample.namespace), _excel_text(sample.pod_name),
            sample.sampled_at, sample.cpu_cores, sample.memory_bytes,
            sample.allocated_cpu_request_cores, sample.allocated_cpu_limit_cores,
            sample.allocated_memory_request_bytes, sample.allocated_memory_limit_bytes,
        ])
        rows_in_sheet += 1

    output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
    try:
        workbook.save(output)
        output.seek(0)
    except Exception:
        output.close()
        raise

    def file_iterator():
        while chunk := output.read(64 * 1024):
            yield chunk

    return StreamingResponse(
        file_iterator(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="resource-analysis-{run_id}.xlsx"'},
        background=BackgroundTask(output.close),
    )


@app.get(
    "/v1/collections/{run_id}/raw-samples",
    response_model=RawUsageSamplesPage,
    response_model_by_alias=True,
)
def collection_raw_samples(
    run_id: str,
    workload_name: Annotated[str, Query(alias="workloadName", min_length=1, max_length=253)],
    db: Annotated[Session, Depends(get_db)],
    authorization: Annotated[str | None, Header()] = None,
    offset: Annotated[int, Query(ge=0)] = 0,
    limit: Annotated[int, Query(ge=1, le=500)] = 100,
):
    authorize(authorization)
    run = get_run(run_id, db)
    if run.status != "completed":
        raise HTTPException(status_code=409, detail=f"Collection is {run.status}")
    query = select(RawUsageSample).where(
        RawUsageSample.run_id == run_id,
        RawUsageSample.workload_name == workload_name,
    )
    total = db.scalar(select(func.count()).select_from(RawUsageSample).where(
        RawUsageSample.run_id == run_id,
        RawUsageSample.workload_name == workload_name,
    ))
    if not total:
        raise HTTPException(status_code=404, detail="No raw samples found for this workload in the collection")
    rows = db.scalars(
        query.order_by(RawUsageSample.sampled_at.desc(), RawUsageSample.pod_name)
        .offset(offset)
        .limit(limit)
    ).all()
    workload_type = db.scalar(
        select(RawUsageSample.workload_type)
        .where(
            RawUsageSample.run_id == run_id,
            RawUsageSample.workload_name == workload_name,
        )
        .limit(1)
    )
    return RawUsageSamplesPage(
        workload_name=workload_name,
        workload_type=workload_type,
        total=total,
        offset=offset,
        limit=limit,
        items=rows,
    )
